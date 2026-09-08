"""Convert the SysML EV requirements/relationships workbook into data/model.json.

Reads the source workbook (Requirements, Requirement_Relationships and
All_Relationships sheets) plus the hand-authored data/blocks.json catalog and
emits a single deterministic data/model.json following docs/model-contract.md
section 2. Uses openpyxl only (no pandas).

Usage:
    python tools/xlsx_to_json.py <xlsx> <blocks.json> <out.json>
    python tools/xlsx_to_json.py --inspect <xlsx>
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import sys
from pathlib import Path

import openpyxl

# ---------------------------------------------------------------------------
# Column maps: logical field name -> workbook header text, one dict per sheet.
# Keeping these here (rather than hardcoding column indexes) means a reordered
# column in the workbook does not silently break the converter.
# ---------------------------------------------------------------------------
COLUMN_MAP = {
    "Requirements": {
        "import_order": "Import Order",
        "external_id": "External ID",
        "requirement_id": "Requirement ID",
        "name": "Name",
        "category": "Category",
        "authoritative": "Authoritative Master",
        "owner_id": "Owner External ID",
        "owner_name": "Owner Name",
        "text": "Requirement Text",
        "rationale": "Rationale",
        "acceptance": "Acceptance Basis",
        "documentation": "Documentation",
    },
    "Requirement_Relationships": {
        "import_order": "Import Order",
        "external_id": "External ID",
        "type": "Relationship Type",
        "source_id": "Source External ID",
        "source_kind": "Source Kind",
        "source_name": "Source Name",
        "target_id": "Target External ID",
        "target_kind": "Target Kind",
        "target_name": "Target Name",
        "owner_id": "Owner External ID",
        "owner_name": "Owner Name",
        "documentation": "Documentation",
    },
    "All_Relationships": {
        "import_order": "Import Order",
        "external_id": "External ID",
        "operation": "Operation",
        "type": "Relationship Type",
        "source_id": "Source / Role",
        "source_kind": "Source Kind",
        "source_name": "Source Name",
        "source_path": "Source Path",
        "target_id": "Target / Role",
        "target_kind": "Target Kind",
        "target_name": "Target Name",
        "target_path": "Target Path",
        "owner_id": "Owner / Context",
        "owner_name": "Owner / Context Name",
        "connector": "Connector",
        "conveyed_items": "Conveyed Items",
        "name": "Name",
        "documentation": "Documentation",
    },
}

# ---------------------------------------------------------------------------
# Port id prefix -> owning block id.
#
# The workbook has no explicit "port owner" column. Port-level Assembly /
# Delegation / ItemFlow rows only give the port id (e.g. CTRL_TORQUE,
# PT_TORQUE); the owning block has to be derived from the id's naming
# convention. This map was reverse-engineered from every port id actually
# present in the workbook (verified: all 26 port ids match exactly one
# prefix below). Order does not matter since no prefix is a prefix of
# another prefix in this set.
# ---------------------------------------------------------------------------
PORT_OWNER_PREFIX = {
    "CTRL_": "VCONTROL",
    "PT_": "POWERTRAIN",
    "INV_": "INVERTER",
    "ENERGY_": "ENERGY",
    "THERM_": "THERMAL",
    "SENS_": "SENSORS",
    "HMI_": "HMI",
    "CHARGE_": "CHARGE",
    "DIAG_": "DIAG",
    "BRAKE_": "BRAKES",
    "MOTOR_": "POWERTRAIN",
    "VEH_": "VEH",
}

# Block ids that appear in the workbook (kind Block) but are intentionally
# absent from data/blocks.json because they are abstract SysML
# generalization targets rather than clickable car parts.
KNOWN_ABSTRACT_BLOCKS = {"VEH_SUBSYSTEM"}

# Relationship types whose Source/Target are port ids that must be lifted to
# their owning block id for the top-level source/target fields.
PORT_LEVEL_TYPES = {"Assembly", "Delegation", "ItemFlow"}

# ---------------------------------------------------------------------------
# Requirement category catalog. `name` is the normalized (ASCII-dash) form;
# the workbook's raw category strings are normalized the same way before
# lookup. `plain` phrasings for the subsystem/verification categories are a
# judgment call (not specified verbatim in the contract) chosen to match the
# style of the STK/SYS/PERF/IF examples.
# ---------------------------------------------------------------------------
CATEGORY_DEFS = [
    {"id": "STK", "name": "Stakeholder / Mission", "plain": "What customers and the mission need", "level": 0},
    {"id": "SYS", "name": "System", "plain": "What the car must do", "level": 1},
    {"id": "PERF", "name": "Performance", "plain": "How well it must do it", "level": 1},
    {"id": "IF", "name": "Interface", "plain": "How parts talk to each other", "level": 1},
    {"id": "PT", "name": "Subsystem - Powertrain", "plain": "What the powertrain must do", "level": 2, "block": "POWERTRAIN"},
    {"id": "BAT", "name": "Subsystem - Energy Storage", "plain": "What the battery pack must do", "level": 2, "block": "ENERGY"},
    {"id": "CTRL", "name": "Subsystem - Vehicle Control", "plain": "What the vehicle controller must do", "level": 2, "block": "VCONTROL"},
    {"id": "BRK", "name": "Subsystem - Brake System", "plain": "What the brakes must do", "level": 2, "block": "BRAKES"},
    {"id": "THM", "name": "Subsystem - Thermal Management", "plain": "What the cooling system must do", "level": 2, "block": "THERMAL"},
    {"id": "SNS", "name": "Subsystem - Vehicle Sensors", "plain": "What the sensors must do", "level": 2, "block": "SENSORS"},
    {"id": "HMI", "name": "Subsystem - Driver Interface", "plain": "What the driver interface must do", "level": 2, "block": "HMI"},
    {"id": "CHG", "name": "Subsystem - Charging", "plain": "What the charging system must do", "level": 2, "block": "CHARGE"},
    {"id": "SRV", "name": "Subsystem - Service / Diagnostics", "plain": "What service and diagnostics must do", "level": 2, "block": "DIAG"},
    {"id": "VER", "name": "Verification View Copy", "plain": "Copies kept for the verification view", "level": 3},
]

_DASH_CHARS = ("‐", "‑", "‒", "–", "—", "―", "−")


def normalize_dashes(text: str) -> str:
    for ch in _DASH_CHARS:
        text = text.replace(ch, "-")
    return text


def clean(value):
    """Strip whitespace from string cells; pass everything else through."""
    if isinstance(value, str):
        value = value.strip()
    return value


def read_sheet_records(ws, colmap):
    """Read a worksheet into a list of dicts keyed by COLUMN_MAP logical names."""
    header_row = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
    header_index = {h: i for i, h in enumerate(header_row) if h is not None}
    missing = [header for header in colmap.values() if header not in header_index]
    if missing:
        raise SystemExit(f"Sheet {ws.title!r} is missing expected column(s): {missing}")
    field_index = {field: header_index[header] for field, header in colmap.items()}
    records = []
    for row in ws.iter_rows(min_row=2, values_only=True):
        if row[field_index["import_order"]] is None:
            continue
        record = {field: clean(row[idx]) for field, idx in field_index.items()}
        records.append(record)
    return records


def extract_item_name(conveyed: str) -> str:
    """'DATA_TORQUE (TorqueCommandData)' -> 'TorqueCommandData'."""
    match = re.search(r"\(([^)]+)\)", conveyed or "")
    return match.group(1) if match else (conveyed or "")


def item_label(item_type: str) -> str:
    """'TorqueCommandData' -> 'Torque command' (drop trailing 'Data' word)."""
    words = re.findall(r"[A-Z][a-z0-9]*|[a-z0-9]+", item_type or "")
    if words and words[-1] == "Data":
        words = words[:-1]
    if not words:
        return item_type
    return " ".join([words[0]] + [w.lower() for w in words[1:]])


def lift_port(port_id: str) -> str:
    for prefix, block in PORT_OWNER_PREFIX.items():
        if port_id.startswith(prefix):
            return block
    raise SystemExit(
        f"Cannot determine owning block for port id {port_id!r}; "
        "add its prefix to PORT_OWNER_PREFIX in tools/xlsx_to_json.py"
    )


def build_category_lookup():
    by_name = {}
    for cat in CATEGORY_DEFS:
        by_name[cat["name"]] = cat
    return by_name


def convert(xlsx_path: Path, blocks_path: Path, out_path: Path) -> dict:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)

    for sheet_name in COLUMN_MAP:
        if sheet_name not in wb.sheetnames:
            raise SystemExit(f"Workbook is missing expected sheet {sheet_name!r}")

    blocks_catalog = json.loads(blocks_path.read_text(encoding="utf-8"))
    blocks_by_id = {b["id"]: b for b in blocks_catalog}

    # -------------------------------------------------------------- Requirements
    req_records = read_sheet_records(wb["Requirements"], COLUMN_MAP["Requirements"])
    category_lookup = build_category_lookup()

    requirements_by_id = {}
    for rec in req_records:
        cat_raw = normalize_dashes(rec["category"])
        cat_def = category_lookup.get(cat_raw)
        if cat_def is None:
            raise SystemExit(
                f"Requirement {rec['external_id']!r} has unknown category {rec['category']!r} "
                f"(normalized: {cat_raw!r}); update CATEGORY_DEFS."
            )
        elem = {
            "id": rec["external_id"],
            "kind": "Requirement",
            "displayId": rec["requirement_id"],
            "category": cat_def["id"],
            "name": rec["name"],
            "text": rec["text"],
            "rationale": rec["rationale"],
            "acceptance": rec["acceptance"],
            "authoritative": rec["authoritative"] == "Yes",
            "owner": rec["owner_id"],
        }
        if elem["id"] in requirements_by_id:
            raise SystemExit(f"Duplicate requirement External ID {elem['id']!r}")
        requirements_by_id[elem["id"]] = elem

    # ------------------------------------------------- Requirement_Relationships
    rr_records = read_sheet_records(
        wb["Requirement_Relationships"], COLUMN_MAP["Requirement_Relationships"]
    )

    # ------------------------------------------------------------- All_Relationships
    ar_records = read_sheet_records(wb["All_Relationships"], COLUMN_MAP["All_Relationships"])

    # Sanity check called out in the task: Requirement_Relationships must be a
    # strict subset of All_Relationships (same external id -> same type/source/target).
    ar_by_id = {rec["external_id"]: rec for rec in ar_records}
    for rec in rr_records:
        counterpart = ar_by_id.get(rec["external_id"])
        if counterpart is None:
            raise SystemExit(
                f"Requirement_Relationships row {rec['external_id']!r} does not appear "
                "in All_Relationships; sheets are expected to be consistent."
            )
        if (
            counterpart["type"] != rec["type"]
            or counterpart["source_id"] != rec["source_id"]
            or counterpart["target_id"] != rec["target_id"]
        ):
            raise SystemExit(
                f"Requirement_Relationships row {rec['external_id']!r} disagrees with "
                "its All_Relationships counterpart."
            )

    # Connector id -> its own relationship type (Assembly rows and Delegation
    # rows *are* connectors; ItemFlow rows merely reference one via `connector`).
    connector_type = {
        rec["external_id"]: rec["type"]
        for rec in ar_records
        if rec["type"] in ("Assembly", "Delegation")
    }

    # Generic name/kind harvesting: every id that appears as a Source or
    # Target anywhere in All_Relationships becomes a candidate element.
    name_of = {}
    kind_of = {}
    for rec in ar_records:
        for id_field, kind_field, name_field in (
            ("source_id", "source_kind", "source_name"),
            ("target_id", "target_kind", "target_name"),
        ):
            elem_id, elem_kind, elem_name = rec[id_field], rec[kind_field], rec[name_field]
            if elem_id is None:
                continue
            if elem_id in kind_of and kind_of[elem_id] != elem_kind:
                raise SystemExit(
                    f"Id {elem_id!r} appears with two different kinds: "
                    f"{kind_of[elem_id]!r} and {elem_kind!r}"
                )
            kind_of[elem_id] = elem_kind
            name_of[elem_id] = elem_name

    # ------------------------------------------------------------------- relationships
    relationships = []
    flows_by_connector = {}
    for rec in ar_records:
        entry = {"id": rec["external_id"], "type": rec["type"]}
        if rec["type"] in PORT_LEVEL_TYPES:
            source_block = lift_port(rec["source_id"])
            target_block = lift_port(rec["target_id"])
            entry["source"] = source_block
            entry["target"] = target_block
            extra = {"sourcePort": rec["source_id"], "targetPort": rec["target_id"]}
            if rec["type"] == "ItemFlow":
                item = extract_item_name(rec["conveyed_items"])
                extra["item"] = item
                extra["connector"] = rec["connector"]
                if connector_type.get(rec["connector"]) == "Assembly":
                    flows_by_connector[rec["connector"]] = {
                        "id": rec["connector"],
                        "source": source_block,
                        "target": target_block,
                        "item": item,
                        "label": item_label(item),
                        "meshName": f"FLOW__{source_block}__{target_block}",
                    }
            entry["extra"] = extra
        else:
            entry["source"] = rec["source_id"]
            entry["target"] = rec["target_id"]
        relationships.append(entry)

    if len(relationships) != len(ar_records):
        raise SystemExit("Internal error: relationship count does not match row count")

    # Attach copyOf to verification-view copy requirements.
    for rec in ar_records:
        if rec["type"] == "Copy":
            copy_id, master_id = rec["source_id"], rec["target_id"]
            if copy_id not in requirements_by_id:
                raise SystemExit(f"Copy relationship {rec['external_id']!r} source {copy_id!r} is not a Requirement")
            requirements_by_id[copy_id]["copyOf"] = master_id

    # ------------------------------------------------------------------------- elements
    elements_by_id = dict(requirements_by_id)

    catalog_block_ids_seen = set()
    for elem_id, raw_kind in kind_of.items():
        if raw_kind == "Requirement":
            continue  # already built from the Requirements sheet
        name = name_of[elem_id]
        if raw_kind == "Block":
            if elem_id in blocks_by_id:
                b = blocks_by_id[elem_id]
                elements_by_id[elem_id] = {
                    "id": elem_id,
                    "kind": "Block",
                    "name": b["name"],
                    "label": b["label"],
                    "blurb": b["blurb"],
                    "category": b["category"],
                    "mesh": b["meshName"],
                    "parent": b["parentId"],
                    "color": b["color"],
                    "alpha": b["alpha"],
                    "explode": b["explode"],
                }
                catalog_block_ids_seen.add(elem_id)
            elif elem_id in KNOWN_ABSTRACT_BLOCKS:
                elements_by_id[elem_id] = {
                    "id": elem_id,
                    "kind": "Block",
                    "name": name,
                    "abstract": True,
                }
            else:
                raise SystemExit(
                    f"Block {elem_id!r} appears in the workbook but is missing from "
                    "data/blocks.json (and is not a known abstract block). "
                    "Add it to blocks.json or KNOWN_ABSTRACT_BLOCKS."
                )
        elif raw_kind in ("ProxyPort", "FullPort"):
            elements_by_id[elem_id] = {
                "id": elem_id,
                "kind": "Port",
                "name": name,
                "owner": lift_port(elem_id),
                "portKind": raw_kind,
            }
        elif raw_kind == "Binding Endpoint":
            elements_by_id[elem_id] = {"id": elem_id, "kind": "BindingEndpoint", "name": name}
        elif raw_kind == "ModelLibrary":
            elements_by_id[elem_id] = {"id": elem_id, "kind": "Package", "name": name}
        else:
            # Package, TestCase, UseCase, Actor, ConstraintBlock, Operation, ...
            elements_by_id[elem_id] = {"id": elem_id, "kind": raw_kind, "name": name}

    missing_catalog_blocks = set(blocks_by_id) - catalog_block_ids_seen
    if missing_catalog_blocks:
        raise SystemExit(
            "blocks.json id(s) never appear in the workbook: "
            f"{sorted(missing_catalog_blocks)}"
        )

    elements = sorted(elements_by_id.values(), key=lambda e: e["id"])
    relationships_sorted = sorted(relationships, key=lambda r: r["id"])
    flows = sorted(flows_by_connector.values(), key=lambda f: f["id"])

    categories = sorted(CATEGORY_DEFS, key=lambda c: (c["level"], c["id"]))

    hierarchy = {}
    for b in blocks_catalog:
        if b["parentId"] is not None:
            hierarchy.setdefault(b["parentId"], []).append(b["id"])
    for parent in hierarchy:
        hierarchy[parent].sort()

    stats = {
        "requirements": sum(1 for r in requirements_by_id.values() if "copyOf" not in r),
        "copies": sum(1 for r in requirements_by_id.values() if "copyOf" in r),
        "traceRelationships": len(rr_records),
        "allRelationships": len(ar_records),
        "blocks": len(blocks_catalog),
        "flows": len(flows),
    }

    source_bytes = xlsx_path.read_bytes()
    source_sha256 = hashlib.sha256(source_bytes).hexdigest()

    model = {
        "meta": {
            "schema": 1,
            "sourceFile": _repo_relative(xlsx_path),
            "sourceSha256": source_sha256,
            "converter": "tools/xlsx_to_json.py",
        },
        "categories": categories,
        "elements": elements,
        "relationships": relationships_sorted,
        "flows": flows,
        "hierarchy": hierarchy,
        "stats": stats,
    }
    return model


def write_model(model: dict, out_path: Path) -> None:
    text = json.dumps(model, indent=2, ensure_ascii=False, sort_keys=False) + "\n"
    out_path.write_text(text, encoding="utf-8", newline="\n")


def inspect_workbook(xlsx_path: Path) -> None:
    wb = openpyxl.load_workbook(xlsx_path, data_only=True)
    print(f"Sheets: {wb.sheetnames}")
    for name in wb.sheetnames:
        ws = wb[name]
        headers = next(ws.iter_rows(min_row=1, max_row=1, values_only=True))
        data_rows = sum(
            1 for row in ws.iter_rows(min_row=2, values_only=True) if row[0] is not None
        )
        print(f"\n=== {name} ({ws.dimensions}) — {data_rows} data rows ===")
        print(f"headers: {list(headers)}")



def _repo_relative(path) -> str:
    """Path relative to the repo root (parent of tools/), posix style, so output is machine-independent."""
    repo_root = Path(__file__).resolve().parent.parent
    p = Path(path).resolve()
    try:
        return p.relative_to(repo_root).as_posix()
    except ValueError:
        return f"data/source/{p.name}"

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--inspect", metavar="XLSX", help="print sheets/headers/row counts and exit")
    parser.add_argument("xlsx", nargs="?", help="input workbook path")
    parser.add_argument("blocks", nargs="?", help="data/blocks.json path")
    parser.add_argument("out", nargs="?", help="output model.json path")
    args = parser.parse_args(argv)

    if args.inspect:
        inspect_workbook(Path(args.inspect))
        return 0

    if not (args.xlsx and args.blocks and args.out):
        parser.error("xlsx, blocks and out are required unless --inspect is given")

    model = convert(Path(args.xlsx), Path(args.blocks), Path(args.out))
    write_model(model, Path(args.out))
    print(
        f"Wrote {args.out}: {model['stats']}"
    )
    return 0


if __name__ == "__main__":
    sys.exit(main())
